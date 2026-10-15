import csv
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden
from django.db import models
from django.utils import timezone

from exam.models import AssessmentAttempt, AttemptResponse
from assessments.models import Assessment, Question
from accounts.models import Department, User, Team
from django.core.paginator import Paginator


@login_required
def results_list_view(request):
    if not (request.user.is_team_lead or request.user.is_trainer or request.user.is_super_admin):
        return HttpResponseForbidden("Access restricted to Supervisors, Trainers, and Administrators.")

    q = request.GET.get('q', '').strip()
    assess_id = request.GET.get('assess', '')
    dept_id = request.GET.get('dept', '')
    team_id = request.GET.get('team', '')
    status_filter = request.GET.get('status', '')
    result_filter = request.GET.get('result', '')

    attempts = AssessmentAttempt.objects.select_related(
        'user__department', 'user__team', 'assessment', 'certificate', 'link'
    ).order_by('-started_at')

    if q:
        attempts = attempts.filter(
            models.Q(user__first_name__icontains=q) |
            models.Q(user__last_name__icontains=q) |
            models.Q(user__username__icontains=q) |
            models.Q(user__employee_id__icontains=q) |
            models.Q(participant_name__icontains=q) |
            models.Q(participant_email__icontains=q) |
            models.Q(participant_employee_id__icontains=q) |
            models.Q(link__name__icontains=q) |
            models.Q(assessment__title__icontains=q) |
            models.Q(assessment__code__icontains=q)
        )
    if assess_id:
        attempts = attempts.filter(assessment_id=assess_id)
    if dept_id:
        attempts = attempts.filter(user__department_id=dept_id)
    if team_id:
        attempts = attempts.filter(
            models.Q(user__team_id=team_id) | models.Q(user__groups_assigned__id=team_id)
        ).distinct()
    if status_filter:
        attempts = attempts.filter(status=status_filter)
    if result_filter == 'pass':
        attempts = attempts.filter(is_passed=True)
    elif result_filter == 'fail':
        attempts = attempts.filter(is_passed=False)

    total_count = attempts.count()
    graded_qs = attempts.filter(status=AssessmentAttempt.Status.GRADED)
    graded_count = graded_qs.count()
    pass_count = graded_qs.filter(is_passed=True).count()
    pass_rate = round((pass_count / graded_count * 100), 1) if graded_count > 0 else 0
    avg_score = graded_qs.aggregate(models.Avg('percentage'))['percentage__avg']
    avg_score = round(avg_score, 1) if avg_score else 0

    paginator = Paginator(attempts, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    departments = Department.objects.filter(is_active=True)
    teams = Team.objects.filter(is_active=True)
    assessments = Assessment.objects.all()

    return render(request, 'analytics/results_list.html', {
        'page_obj': page_obj,
        'total_count': total_count,
        'graded_count': graded_count,
        'pass_rate': pass_rate,
        'avg_score': avg_score,
        'departments': departments,
        'teams': teams,
        'assessments': assessments,
        'q': q,
        'assess_id': assess_id,
        'dept_id': dept_id,
        'team_id': team_id,
        'status_filter': status_filter,
        'result_filter': result_filter,
        'statuses': AssessmentAttempt.Status.choices,
    })


@login_required
def reports_overview_view(request):
    if not (request.user.is_team_lead or request.user.is_trainer or request.user.is_super_admin):
        return HttpResponseForbidden("Access restricted to Supervisors, Trainers, and Administrators.")

    dept_filter = request.GET.get('dept', '')
    assess_filter = request.GET.get('assess', '')

    attempts = AssessmentAttempt.objects.select_related('user__department', 'assessment').all()

    if dept_filter:
        attempts = attempts.filter(user__department_id=dept_filter)
    if assess_filter:
        attempts = attempts.filter(assessment_id=assess_filter)

    total_attempts = attempts.count()
    graded_attempts = attempts.filter(status=AssessmentAttempt.Status.GRADED)
    graded_count = graded_attempts.count()

    passed_count = graded_attempts.filter(is_passed=True).count()
    pass_rate = round((passed_count / graded_count * 100), 1) if graded_count > 0 else 0

    avg_score = graded_attempts.aggregate(models.Avg('percentage'))['percentage__avg']
    avg_score = round(avg_score, 1) if avg_score else 0

    departments = Department.objects.filter(is_active=True)
    assessments = Assessment.objects.all()

    # Department Breakdown
    dept_stats = []
    for d in departments:
        d_attempts = graded_attempts.filter(user__department=d)
        d_count = d_attempts.count()
        d_pass = d_attempts.filter(is_passed=True).count()
        d_avg = d_attempts.aggregate(models.Avg('percentage'))['percentage__avg']
        dept_stats.append({
            'department': d,
            'count': d_count,
            'pass_rate': round((d_pass / d_count * 100), 1) if d_count > 0 else 0,
            'avg_score': round(d_avg, 1) if d_avg else 0,
        })

    # Commonly Missed Questions
    missed_responses = AttemptResponse.objects.filter(
        attempt__in=graded_attempts,
        marks_awarded=0
    ).values(
        'snapshot__question_text',
        'snapshot__question_type'
    ).annotate(
        miss_count=models.Count('id')
    ).order_by('-miss_count')[:6]

    return render(request, 'analytics/reports.html', {
        'total_attempts': total_attempts,
        'graded_count': graded_count,
        'pass_rate': pass_rate,
        'avg_score': avg_score,
        'dept_stats': dept_stats,
        'missed_responses': missed_responses,
        'departments': departments,
        'assessments': assessments,
        'dept_filter': dept_filter,
        'assess_filter': assess_filter,
        'recent_attempts': attempts.order_by('-started_at')[:15],
    })


@login_required
def export_attempts_csv_view(request):
    if not (request.user.is_trainer or request.user.is_super_admin):
        return HttpResponseForbidden("Access restricted.")

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="breakthru_assessment_attempts.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Attempt ID', 'Candidate Name', 'Employee ID', 'Department',
        'Link Source', 'Assessment Code', 'Assessment Title', 'Version', 'Attempt Number',
        'Status', 'Started At', 'Submitted At', 'Earned Marks', 'Max Marks',
        'Percentage', 'Result'
    ])

    attempts = AssessmentAttempt.objects.select_related('user__department', 'assessment', 'link').order_by('-started_at')

    for att in attempts:
        c_name = att.participant_name or att.user.get_full_name() or att.user.username
        c_emp = att.participant_employee_id or att.user.employee_id or ''
        c_dept = att.participant_department or (att.user.department.name if att.user.department else '')
        link_str = att.link.name if att.link else 'Direct / Group'

        writer.writerow([
            att.id,
            c_name,
            c_emp,
            c_dept,
            link_str,
            att.assessment.code,
            att.assessment.title,
            att.assessment_version,
            att.attempt_number,
            att.status,
            att.started_at.strftime('%Y-%m-%d %H:%M:%S'),
            att.submitted_at.strftime('%Y-%m-%d %H:%M:%S') if att.submitted_at else '',
            att.earned_marks if att.earned_marks is not None else '',
            att.max_possible_marks,
            att.percentage if att.percentage is not None else '',
            'PASSED' if att.is_passed is True else ('FAILED' if att.is_passed is False else 'PENDING')
        ])

    return response


@login_required
def export_attempts_excel_view(request):
    if not (request.user.is_trainer or request.user.is_super_admin):
        return HttpResponseForbidden("Access restricted.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment Attempts"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0B192C", end_color="0B192C", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    headers = [
        'Attempt ID', 'Candidate Name', 'Employee ID', 'Department',
        'Link Source', 'Assessment Code', 'Assessment Title', 'Attempt #',
        'Status', 'Score', 'Percentage', 'Result', 'Date'
    ]

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    attempts = AssessmentAttempt.objects.select_related('user__department', 'assessment', 'link').order_by('-started_at')

    for att in attempts:
        result_str = 'PASSED' if att.is_passed is True else ('FAILED' if att.is_passed is False else 'PENDING')
        c_name = att.participant_name or att.user.get_full_name() or att.user.username
        c_emp = att.participant_employee_id or att.user.employee_id or ''
        c_dept = att.participant_department or (att.user.department.name if att.user.department else '')
        link_str = att.link.name if att.link else 'Direct / Group'

        ws.append([
            att.id,
            c_name,
            c_emp,
            c_dept,
            link_str,
            att.assessment.code,
            att.assessment.title,
            att.attempt_number,
            att.status,
            f"{att.earned_marks}/{att.max_possible_marks}" if att.earned_marks is not None else '',
            f"{att.percentage}%" if att.percentage is not None else '',
            result_str,
            att.started_at.strftime('%Y-%m-%d %H:%M')
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="BreakThru_Assessment_Reports.xlsx"'
    return response
